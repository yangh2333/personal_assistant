"""
Excel导出工具
使用openpyxl生成专业的Excel报表
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import json
import os

class ExcelExporter:
    """Excel导出器"""
    
    def __init__(self, export_dir):
        self.export_dir = export_dir
        self._setup_styles()
    
    def _setup_styles(self):
        """设置样式"""
        # 标题行样式
        self.header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        self.header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
        
        # 交替行样式
        self.alt_row_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
        
        # 常规样式
        self.normal_font = Font(name='微软雅黑', size=10)
        self.normal_align = Alignment(horizontal='left', vertical='center')
        self.currency_align = Alignment(horizontal='right', vertical='center')
        
        # 边框
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def export(self, calculation):
        """
        导出计算结果为Excel文件
        
        Args:
            calculation: Calculation模型实例
            
        Returns:
            str: 导出文件的完整路径
        """
        wb = Workbook()
        
        # 移除默认sheet
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']
        
        # 获取计算详情
        details = json.loads(calculation.calculation_details) if calculation.calculation_details else {}
        
        # 创建各工作表
        self._create_summary_sheet(wb, calculation, details)
        self._create_hardware_sheet(wb, details)
        self._create_software_sheet(wb, details)
        self._create_ops_sheet(wb, details)
        self._create_params_sheet(wb, details)
        
        # 保存文件
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        file_path = os.path.join(self.export_dir, f'ai_cost_report_{timestamp}.xlsx')
        wb.save(file_path)
        
        return file_path
    
    def _create_summary_sheet(self, wb, calculation, details):
        """创建汇总表"""
        ws = wb.create_sheet('汇总表', 0)
        
        # 设置列宽
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 15
        
        # 标题
        ws.merge_cells('A1:C1')
        ws['A1'] = 'AI集群成本监控报告'
        ws['A1'].font = Font(name='微软雅黑', size=16, bold=True)
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30
        
        # 报告生成时间
        ws['A2'] = '生成时间'
        ws['B2'] = calculation.calculated_at.strftime('%Y-%m-%d %H:%M:%S')
        
        # 汇总数据
        row = 4
        ws[f'A{row}'] = '成本汇总'
        ws[f'A{row}'].font = Font(name='微软雅黑', size=12, bold=True)
        row += 1
        
        summary_data = [
            ('总投资额', details.get('total_cost', 0), '元'),
            ('硬件总成本', details.get('total_hardware_cost', 0), '元'),
            ('软件总成本(5年)', details.get('total_software_cost', 0) * 5, '元'),
            ('运维总成本', details.get('total_ops_cost', 0), '元'),
            ('年运营成本', details.get('annual_opex', 0), '元/年'),
            ('5年TCO', details.get('five_year_tco', 0), '元'),
            ('单卡成本', details.get('cost_per_gpu', 0), '元/卡')
        ]
        
        for label, value, unit in summary_data:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            ws[f'B{row}'].number_format = '¥#,##0.00'
            ws[f'C{row}'] = unit
            ws[f'A{row}'].font = self.normal_font
            ws[f'B{row}'].font = self.normal_font
            ws[f'C{row}'].font = self.normal_font
            row += 1
        
        # 成本构成占比
        row += 1
        ws[f'A{row}'] = '成本构成占比'
        ws[f'A{row}'].font = Font(name='微软雅黑', size=12, bold=True)
        row += 1
        
        total = details.get('total_cost', 1)
        hw_pct = details.get('total_hardware_cost', 0) / total * 100
        sw_pct = details.get('total_software_cost', 0) * 5 / total * 100
        ops_pct = details.get('total_ops_cost', 0) / total * 100
        
        ws[f'A{row}'] = '硬件'
        ws[f'B{row}'] = hw_pct / 100
        ws[f'B{row}'].number_format = '0.0%'
        row += 1
        
        ws[f'A{row}'] = '软件'
        ws[f'B{row}'] = sw_pct / 100
        ws[f'B{row}'].number_format = '0.0%'
        row += 1
        
        ws[f'A{row}'] = '运维'
        ws[f'B{row}'] = ops_pct / 100
        ws[f'B{row}'].number_format = '0.0%'
    
    def _create_hardware_sheet(self, wb, details):
        """创建硬件成本明细表"""
        ws = wb.create_sheet('硬件成本明细')
        
        # 设置列宽
        headers = ['成本项', '类别', '单位', '单价(元)', '数量', '小计(元)', '生命周期成本(元)', '年维护费(元)']
        widths = [25, 15, 10, 15, 12, 18, 20, 15]
        
        for i, (header, width) in enumerate(zip(headers, widths), 1):
            ws.column_dimensions[get_column_letter(i)].width = width
            cell = ws.cell(row=1, column=i)
            cell.value = header
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        ws.row_dimensions[1].height = 25
        
        # 数据行
        hardware_items = details.get('hardware', [])
        for idx, item in enumerate(hardware_items, 2):
            is_alt = idx % 2 == 0
            fill = self.alt_row_fill if is_alt else None
            
            data = [
                item.get('item_name', ''),
                item.get('sub_category', ''),
                item.get('unit', ''),
                item.get('unit_price', 0),
                item.get('quantity', 0),
                item.get('subtotal', 0),
                item.get('lifetime_cost', 0),
                item.get('subtotal', 0) * item.get('annual_maintenance_rate', 0) if item.get('annual_maintenance_rate', 0) else 0
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=idx, column=col)
                cell.value = value
                cell.font = self.normal_font
                cell.border = self.thin_border
                
                if fill:
                    cell.fill = fill
                
                if col in [4, 5, 6, 7, 8]:
                    cell.alignment = self.currency_align
                    cell.number_format = '#,##0.00' if col > 4 else '#,##0'
        
        # 添加合计行
        total_row = len(hardware_items) + 2
        ws.cell(row=total_row, column=1).value = '合计'
        ws.cell(row=total_row, column=1).font = Font(name='微软雅黑', size=10, bold=True)
        ws.cell(row=total_row, column=7).value = details.get('total_hardware_cost', 0)
        ws.cell(row=total_row, column=7).number_format = '¥#,##0.00'
        ws.cell(row=total_row, column=7).font = Font(name='微软雅黑', size=10, bold=True)
    
    def _create_software_sheet(self, wb, details):
        """创建软件成本明细表"""
        ws = wb.create_sheet('软件成本明细')
        
        headers = ['成本项', '类别', '单位', '单价(元/年)', '数量', '年费(元)', '5年费用(元)']
        widths = [25, 15, 12, 15, 12, 15, 15]
        
        for i, (header, width) in enumerate(zip(headers, widths), 1):
            ws.column_dimensions[get_column_letter(i)].width = width
            cell = ws.cell(row=1, column=i)
            cell.value = header
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        ws.row_dimensions[1].height = 25
        
        software_items = details.get('software', [])
        for idx, item in enumerate(software_items, 2):
            is_alt = idx % 2 == 0
            fill = self.alt_row_fill if is_alt else None
            
            annual = item.get('subtotal', 0)
            
            data = [
                item.get('item_name', ''),
                item.get('sub_category', ''),
                item.get('unit', ''),
                item.get('unit_price', 0),
                item.get('quantity', 0),
                annual,
                annual * 5
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=idx, column=col)
                cell.value = value
                cell.font = self.normal_font
                cell.border = self.thin_border
                
                if fill:
                    cell.fill = fill
                
                if col in [4, 5, 6, 7]:
                    cell.alignment = self.currency_align
                    cell.number_format = '#,##0.00' if col > 4 else '#,##0'
        
        # 合计行
        total_row = len(software_items) + 2
        ws.cell(row=total_row, column=1).value = '合计'
        ws.cell(row=total_row, column=1).font = Font(name='微软雅黑', size=10, bold=True)
        ws.cell(row=total_row, column=6).value = details.get('total_software_cost', 0)
        ws.cell(row=total_row, column=6).number_format = '¥#,##0.00'
        ws.cell(row=total_row, column=6).font = Font(name='微软雅黑', size=10, bold=True)
        ws.cell(row=total_row, column=7).value = details.get('total_software_cost', 0) * 5
        ws.cell(row=total_row, column=7).number_format = '¥#,##0.00'
        ws.cell(row=total_row, column=7).font = Font(name='微软雅黑', size=10, bold=True)
    
    def _create_ops_sheet(self, wb, details):
        """创建运维成本明细表"""
        ws = wb.create_sheet('运维成本明细')
        
        headers = ['成本项', '类别', '单位', '年费(元)', '5年费用(元)']
        widths = [25, 15, 12, 18, 18]
        
        for i, (header, width) in enumerate(zip(headers, widths), 1):
            ws.column_dimensions[get_column_letter(i)].width = width
            cell = ws.cell(row=1, column=i)
            cell.value = header
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        ws.row_dimensions[1].height = 25
        
        ops_items = details.get('ops', [])
        
        # 添加电力成本
        electricity = details.get('electricity', {})
        if electricity:
            ws.cell(row=2, column=1).value = '电力成本'
            ws.cell(row=2, column=2).value = 'electricity'
            ws.cell(row=2, column=3).value = '年'
            ws.cell(row=2, column=4).value = electricity.get('annual_cost', 0)
            ws.cell(row=2, column=5).value = electricity.get('lifetime_cost', 0)
        
        for idx, item in enumerate(ops_items, 3):
            is_alt = idx % 2 == 0
            fill = self.alt_row_fill if is_alt else None
            
            data = [
                item.get('item_name', ''),
                item.get('sub_category', ''),
                item.get('unit', ''),
                item.get('annual_cost', 0),
                item.get('lifetime_cost', 0)
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=idx, column=col)
                cell.value = value
                cell.font = self.normal_font
                cell.border = self.thin_border
                
                if fill:
                    cell.fill = fill
                
                if col in [4, 5]:
                    cell.alignment = self.currency_align
                    cell.number_format = '¥#,##0.00'
        
        # 合计行
        total_row = len(ops_items) + 3
        ws.cell(row=total_row, column=1).value = '合计'
        ws.cell(row=total_row, column=1).font = Font(name='微软雅黑', size=10, bold=True)
        ws.cell(row=total_row, column=4).value = details.get('total_ops_cost', 0) / 5
        ws.cell(row=total_row, column=4).number_format = '¥#,##0.00'
        ws.cell(row=total_row, column=4).font = Font(name='微软雅黑', size=10, bold=True)
        ws.cell(row=total_row, column=5).value = details.get('total_ops_cost', 0)
        ws.cell(row=total_row, column=5).number_format = '¥#,##0.00'
        ws.cell(row=total_row, column=5).font = Font(name='微软雅黑', size=10, bold=True)
    
    def _create_params_sheet(self, wb, details):
        """创建计算参数表"""
        ws = wb.create_sheet('计算参数')
        
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 15
        
        headers = ['参数名称', '参数值', '单位']
        for i, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=i)
            cell.value = header
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        ws.row_dimensions[1].height = 25
        
        params = [
            ('GPU数量', details.get('gpu_count', 0), '卡'),
            ('GPU型号', details.get('gpu_model', 'N/A'), ''),
            ('服务器数量', details.get('server_count', 0), '台'),
            ('功耗', details.get('power_consumption_kw', 0), 'KW'),
            ('电价', details.get('electricity_price', 0.6), '元/度'),
            ('PUE值', details.get('pue_ratio', 1.3), ''),
            ('部署方式', details.get('location_type', 'colocation'), '')
        ]
        
        for idx, (name, value, unit) in enumerate(params, 2):
            ws.cell(row=idx, column=1).value = name
            ws.cell(row=idx, column=2).value = value
            ws.cell(row=idx, column=3).value = unit
            
            for col in range(1, 4):
                ws.cell(row=idx, column=col).font = self.normal_font
                ws.cell(row=idx, column=col).border = self.thin_border
